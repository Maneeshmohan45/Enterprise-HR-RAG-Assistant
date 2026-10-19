"""File loaders: turn each supported file type into LangChain ``Document`` objects.

Design rules
------------
* Prose files (PDF, DOCX, TXT, HTML) are split into **sections** at headings,
  so every Document is one logical section (e.g. "3. Leave Categories").
* Record files (CSV, JSON) produce **one Document per record** (one employee).
* Excel produces **one Document per sheet**, with one line per row.
* Every Document gets ``file_name``, ``file_type`` and ``content_kind``
  ("prose" or "records"). Where it makes sense it also gets ``section_title``,
  ``page_number``, ``row_number``, ``sheet_name`` or ``employee_id``.

No heavy dependencies: pypdf, python-docx, openpyxl, beautifulsoup4 only.
"""
from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Iterable

from langchain_core.documents import Document

# --------------------------------------------------------------------------
# Section splitting (shared by PDF / DOCX / TXT / HTML)
# --------------------------------------------------------------------------
_MD_HEADING = re.compile(r"^#\s+(\S.*?)\s*$")
# Whole-line numbered heading: "3. Leave Categories"  (not "1. 1. Table of contents")
_NUM_HEADING = re.compile(r"^\d{1,2}\.\s+(?!\d+\.)([A-Z][^\n]{1,70}?)\s*$")
# Heading + text on one line, heading in capitals: "5.  ATTENDANCE Employees must:"
_NUM_CAPS_INLINE = re.compile(
    r"^\d{1,2}\.\s+([A-Z][A-Z0-9\-/&]+(?:\s+[A-Z][A-Z0-9\-/&]+)*)(?:\s+(.*))?$"
)


def _tidy_title(title: str) -> str:
    """Collapse repeated spaces and turn ALL-CAPS headings into Title Case."""
    title = re.sub(r"\s+", " ", title).strip()
    return title.title() if title.isupper() else title


def detect_heading(line: str) -> tuple[str, str] | None:
    """Return ``(heading, remainder_of_line)`` if ``line`` starts a section."""
    line = line.strip()
    if not line:
        return None
    m = _MD_HEADING.match(line)
    if m:
        return _tidy_title(m.group(1)), ""

    # "5.  ATTENDANCE Employees must:" -> heading "5. Attendance" + body text.
    # Needs 4+ capital letters so "10. HR Policies" is not split after "HR".
    m = _NUM_CAPS_INLINE.match(line)
    if m and len(re.sub(r"[^A-Z]", "", m.group(1))) >= 4:
        end = line.index(m.group(1)) + len(m.group(1))
        return _tidy_title(line[:end]), (m.group(2) or "")

    # "3. Leave Categories" (whole line is the heading)
    m = _NUM_HEADING.match(line)
    if m:
        title = m.group(1)
        # Real headings are short and are not sentences or list intros.
        if len(title.split()) <= 9 and not title.endswith((".", ":", ",", ";")):
            return _tidy_title(line), ""
    return None


def split_sections(
    lines: Iterable[tuple[int | None, str]], default_title: str = "Overview"
) -> list[tuple[str, int | None, str]]:
    """Group ``(page, line)`` pairs into ``(title, start_page, text)`` sections."""
    sections: list[tuple[str, int | None, str]] = []
    title, page, buf = default_title, None, []
    started = False

    def flush():
        text = "\n".join(buf).strip()
        if text:
            sections.append((title, page, text))

    for line_page, line in lines:
        if not started and page is None:
            page = line_page
        heading = detect_heading(line)
        if heading:
            flush()
            title, page, buf = heading[0], line_page, []
            started = True
            if heading[1]:
                buf.append(heading[1])
        else:
            buf.append(line)
    flush()
    return sections


def _sections_to_documents(
    sections: list[tuple[str, int | None, str]], base: dict
) -> list[Document]:
    docs = []
    for index, (title, page, text) in enumerate(sections, start=1):
        meta = {**base, "section_title": title, "section_index": index}
        if page is not None:
            meta["page_number"] = page
        docs.append(Document(page_content=text, metadata=meta))
    return docs


# --------------------------------------------------------------------------
# Table helper (DOCX + HTML)
# --------------------------------------------------------------------------
def table_to_lines(rows: list[list[str]]) -> list[str]:
    """Turn table rows into self-describing lines.

    A bare row like ``Annual Leave | 18 Days | 20 Days`` loses its meaning when
    the header row is dropped, so we keep the column names:
    ``Leave Type: Annual Leave | V1: 18 Days | V2: 20 Days``.
    """
    rows = [[c.strip() for c in row] for row in rows if any(c.strip() for c in row)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    lines: list[str] = []
    if width == 2:  # key / value table
        for key, *rest in rows:
            value = rest[0] if rest else ""
            if key and value:
                lines.append(f"{key}: {value}")
        return lines
    header, body = rows[0], rows[1:]
    for row in body:
        cells = [f"{h}: {v}" for h, v in zip(header, row) if h and v]
        if cells:
            lines.append(" | ".join(cells))
    return lines


# --------------------------------------------------------------------------
# Dates: make ISO dates searchable by month name
# --------------------------------------------------------------------------
_ISO_DATE = re.compile(r"\b(\d{4})-(\d{2})-(\d{2})\b")


def humanize_dates(text: str) -> str:
    """``2025-01-14`` -> ``2025-01-14 (14 January 2025)``.

    People ask "which holidays are in January?", but spreadsheets store
    ``2025-01-14``. Adding the long form lets both keyword and semantic search
    match. Impossible dates (e.g. ``1995-31-12``) are left untouched.
    """
    def repl(m):
        try:
            d = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return m.group(0)
        return f"{m.group(0)} ({d.day} {d:%B %Y})"

    return _ISO_DATE.sub(repl, text)


# --------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------
def load_txt(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = [(None, ln) for ln in text.splitlines()]
    base = {"file_name": path.name, "file_type": ".txt", "content_kind": "prose"}
    return _sections_to_documents(split_sections(lines), base)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
_PAGE_NUMBER_LINE = re.compile(r"(?i)^\s*(?:page\s*)?\d+(?:\s*of\s*\d+)?\s*$")
# Table-of-contents entries are printed with a doubled number: "1. 1. Welcome"
_TOC_LINE = re.compile(r"^\s*\d{1,2}\.\s+\d{1,2}\.\s+\S")


def strip_repeated_lines(pages: list[list[str]], min_pages: int = 3, ratio: float = 0.6):
    """Remove running headers/footers (a line that repeats on most pages)."""
    if len(pages) < min_pages:
        return pages

    def norm(line: str) -> str:
        return re.sub(r"\d+", "#", line.strip().lower())

    presence = Counter()
    for page in pages:
        presence.update({norm(line) for line in page if line.strip()})
    limit = max(min_pages, ratio * len(pages))
    repeated = {line for line, n in presence.items() if n >= limit}
    return [[ln for ln in page if norm(ln) not in repeated] for page in pages]


def load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").splitlines() for page in reader.pages]
    pages = strip_repeated_lines(pages)
    lines = [
        (page_no, line)
        for page_no, page in enumerate(pages, start=1)
        for line in page
        if not _PAGE_NUMBER_LINE.match(line) and not _TOC_LINE.match(line)
    ]
    base = {"file_name": path.name, "file_type": ".pdf", "content_kind": "prose"}
    return _sections_to_documents(split_sections(lines), base)


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def load_docx(path: Path) -> list[Document]:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = DocxDocument(str(path))
    lines: list[tuple[int | None, str]] = []

    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style == "Heading 1":
                lines.append((None, f"# {text}"))
            elif style.startswith("List"):
                lines.append((None, f"- {text}"))
            else:
                lines.append((None, text))
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells, seen = [], set()
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    cells.append(cell.text.strip())
                rows.append(cells)
            lines.extend((None, ln) for ln in table_to_lines(rows))

    base = {"file_name": path.name, "file_type": ".docx", "content_kind": "prose"}
    return _sections_to_documents(split_sections(lines), base)


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------
def load_html(path: Path) -> list[Document]:
    from bs4 import BeautifulSoup, NavigableString

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "noscript", "head"]):
        tag.decompose()
    for table in soup.find_all("table"):
        rows = [
            [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
            for tr in table.find_all("tr")
        ]
        table.replace_with(NavigableString("\n" + "\n".join(table_to_lines(rows)) + "\n"))
    for h in soup.find_all(["h1", "h2"]):
        h.replace_with(NavigableString(f"\n# {h.get_text(' ', strip=True)}\n"))
    for li in soup.find_all("li"):
        li.replace_with(NavigableString(f"\n- {li.get_text(' ', strip=True)}\n"))

    # Put line breaks around block elements only. Inline tags such as <strong>
    # or <span> must not split a sentence into separate lines.
    for br in soup.find_all("br"):
        br.replace_with(NavigableString("\n"))
    for block in soup.find_all(["p", "div", "section", "article", "header", "footer", "ul", "ol", "dl", "dt", "dd", "blockquote"]):
        block.insert_before(NavigableString("\n"))
        block.insert_after(NavigableString("\n"))
    text = soup.get_text("")
    lines = [(None, re.sub(r"\s+", " ", ln).strip()) for ln in text.splitlines() if ln.strip()]
    base = {"file_name": path.name, "file_type": ".html", "content_kind": "prose"}
    return _sections_to_documents(split_sections(lines), base)


# --------------------------------------------------------------------------
# CSV  (one Document per row)
# --------------------------------------------------------------------------
_NAME_COLUMNS = ("employee_name", "full_name", "name", "title")


def load_csv(path: Path) -> list[Document]:
    docs = []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for row_number, row in enumerate(csv.DictReader(handle), start=1):
            row = {(k or "").strip(): (v or "").strip() for k, v in row.items()}
            if not any(row.values()):
                continue
            text = humanize_dates("\n".join(f"{k}: {v}" for k, v in row.items() if v))
            title = next((row[c] for c in _NAME_COLUMNS if row.get(c)), f"Row {row_number}")
            docs.append(
                Document(
                    page_content=text,
                    metadata={
                        "file_name": path.name,
                        "file_type": ".csv",
                        "content_kind": "records",
                        "section_title": title,
                        "row_number": row_number,
                        "employee_id": row.get("employee_id", ""),
                    },
                )
            )
    return docs


# --------------------------------------------------------------------------
# JSON  (one Document per record)
# --------------------------------------------------------------------------
def _label(key: str) -> str:
    return key.replace("_", " ").strip().title()


def flatten_json(value, path: tuple[str, ...] = ()) -> list[str]:
    """Flatten nested JSON into readable ``Path - To - Field: value`` lines."""
    lines: list[str] = []
    if isinstance(value, dict):
        for key, item in value.items():
            lines += flatten_json(item, path + (_label(str(key)),))
    elif isinstance(value, list):
        if all(not isinstance(i, (dict, list)) for i in value):
            items = [str(i) for i in value if i not in (None, "")]
            if items:
                lines.append(f"{' - '.join(path)}: {', '.join(items)}")
        else:
            for index, item in enumerate(value, start=1):
                lines += flatten_json(item, path + (str(index),))
    elif value not in (None, ""):
        lines.append(f"{' - '.join(path)}: {value}")
    return lines


def _find_first(obj, keys: tuple[str, ...]):
    if isinstance(obj, dict):
        for k in keys:
            if isinstance(obj.get(k), (str, int)) and str(obj[k]).strip():
                return str(obj[k])
        for v in obj.values():
            found = _find_first(v, keys)
            if found:
                return found
    return None


def _extract_records(data) -> list[dict]:
    """Find the list of records inside a JSON file, whatever its wrapper."""
    if isinstance(data, list):
        return [r for r in data if isinstance(r, dict)]
    if isinstance(data, dict):
        for value in data.values():
            if isinstance(value, list) and value and all(isinstance(i, dict) for i in value):
                return value
        for key, value in data.items():
            if key.endswith("_record") and isinstance(value, dict):
                # Keep the wrapper's other fields (record_status, expected_validation_errors, ...)
                # in front of the record, so the text carries the file's own validation notes.
                siblings = {k: v for k, v in data.items() if k != key}
                return [{**siblings, **value}]
        return [data]
    return []


def load_json(path: Path) -> list[Document]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    docs = []
    for index, record in enumerate(_extract_records(data), start=1):
        text = humanize_dates("\n".join(flatten_json(record)))
        title = _find_first(record, ("full_name", "employee_name", "name", "title"))
        docs.append(
            Document(
                page_content=text,
                metadata={
                    "file_name": path.name,
                    "file_type": ".json",
                    "content_kind": "records",
                    "section_title": title or f"Record {index}",
                    "row_number": index,
                    "employee_id": str(record.get("employee_id", "") or ""),
                },
            )
        )
    return docs


# --------------------------------------------------------------------------
# XLSX  (one Document per sheet)
# --------------------------------------------------------------------------
def _format_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_xlsx(path: Path) -> list[Document]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    docs = []
    for sheet in workbook.worksheets:
        rows = [[_format_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue

        # Header = first row with >= 2 filled cells that are all text (not numbers).
        header_index = next(
            (
                i
                for i, r in enumerate(rows)
                if sum(1 for c in r if c) >= 2
                and all(not re.fullmatch(r"[\d.,\-: ]+", c) for c in r if c)
            ),
            None,
        )
        lines: list[str] = []
        if header_index is None:
            lines = [" | ".join(c for c in r if c) for r in rows]
        else:
            lines += [" | ".join(c for c in r if c) for r in rows[:header_index]]
            header = rows[header_index]
            for r in rows[header_index + 1:]:
                filled = [(h, v) for h, v in zip(header, r) if v]
                if len(filled) >= 2:
                    lines.append("; ".join(f"{h or 'Column'}: {v}" for h, v in filled))
                else:
                    lines.append(" | ".join(c for c in r if c))

        docs.append(
            Document(
                page_content=humanize_dates("\n".join(lines)),
                metadata={
                    "file_name": path.name,
                    "file_type": ".xlsx",
                    "content_kind": "records",
                    "section_title": sheet.title,
                    "sheet_name": sheet.title,
                },
            )
        )
    workbook.close()
    return docs


# --------------------------------------------------------------------------
# Dispatcher
# --------------------------------------------------------------------------
LOADERS: dict[str, Callable[[Path], list[Document]]] = {
    ".txt": load_txt,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".csv": load_csv,
    ".json": load_json,
    ".xlsx": load_xlsx,
    ".html": load_html,
}


def load_file(path: Path) -> list[Document]:
    """Load one file with the loader that matches its extension."""
    suffix = path.suffix.lower()
    if suffix not in LOADERS:
        raise ValueError(f"Unsupported file type: {suffix} ({path.name})")
    return LOADERS[suffix](path)
